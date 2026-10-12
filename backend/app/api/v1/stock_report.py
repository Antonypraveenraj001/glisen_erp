from datetime import datetime
from io import BytesIO

from fastapi import (
    APIRouter,
    Depends,
    Query,
)
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.dependencies.database import get_db
from app.dependencies.permissions import require_permission
from app.models.company_settings import CompanySettings
from app.models.user import User
from app.schemas.stock_movement import StockMovementResponse
from app.schemas.stock_report import StockSummaryResponse
from app.services.stock_movement_service import StockMovementService
from app.services.stock_report_service import StockReportService


router = APIRouter(
    prefix="/stock-report",
    tags=["Stock Report"],
)


@router.get(
    "/summary",
    response_model=StockSummaryResponse,
)
def get_stock_summary(
    search: str | None = Query(
        default=None,
        description=(
            "Search by product code, "
            "product name, category or HSN code"
        ),
    ),
    stock_status: str | None = Query(
        default=None,
        description=(
            "Filter by In Stock, Low Stock, "
            "Out of Stock or Over Stock"
        ),
    ),
    db: Session = Depends(
        get_db
    ),
    current_user: User = Depends(
        require_permission(
            "stock.view"
        )
    ),
):
    return (
        StockReportService
        .get_stock_summary(
            db=db,
            search=search,
            stock_status=stock_status,
        )
    )


@router.get(
    "/live-stock/excel",
)
def download_live_stock_excel(
    scope: str = Query(
        default="filtered",
        pattern="^(filtered|whole)$",
    ),
    search: str | None = Query(
        default=None,
    ),
    stock_status: str | None = Query(
        default=None,
    ),
    db: Session = Depends(
        get_db
    ),
    current_user: User = Depends(
        require_permission(
            "stock.view"
        )
    ),
):

    normalized_scope = (
        scope
        .strip()
        .lower()
    )

    summary = (
        StockReportService
        .get_stock_summary(
            db=db,
            search=(
                search
                if normalized_scope
                ==
                "filtered"
                else None
            ),
            stock_status=(
                stock_status
                if normalized_scope
                ==
                "filtered"
                else None
            ),
        )
    )

    live_items = [
        item
        for item
        in summary.items
        if item.current_stock > 0
    ]

    # The ERP is branded CASTLE, but a stock report belongs to
    # the manufacturing company configured in Settings.
    company = db.query(CompanySettings).first()

    def plain(value) -> str:
        return " ".join(str(value or "").split())

    company_name = plain(company.company_name) if company else ""
    address = plain(company.address) if company else ""
    state = plain(company.state_name) if company else ""
    gstin = plain(company.gst_number) if company else ""
    phone = plain(company.phone) if company else ""
    email = plain(company.email) if company else ""

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Live Stock"
    worksheet.freeze_panes = "A6"

    worksheet.merge_cells("A1:H1")
    worksheet["A1"] = company_name or "LIVE STOCK REPORT"
    worksheet["A1"].data_type = "s"
    worksheet["A1"].font = Font(
        bold=True,
        size=16,
    )
    worksheet["A1"].alignment = Alignment(
        horizontal="center",
    )

    worksheet.merge_cells("A2:H2")
    worksheet["A2"] = ", ".join(
        part for part in (address, state) if part
    )
    worksheet["A2"].data_type = "s"
    worksheet["A2"].font = Font(
        size=10,
    )
    worksheet["A2"].alignment = Alignment(
        horizontal="center",
        wrap_text=True,
    )
    worksheet.row_dimensions[2].height = 30

    worksheet.merge_cells("A3:H3")
    worksheet["A3"] = " | ".join(
        part for part in (
            f"GSTIN: {gstin}" if gstin else "",
            f"Phone: {phone}" if phone else "",
            f"Email: {email}" if email else "",
        ) if part
    )
    worksheet["A3"].data_type = "s"
    worksheet["A3"].alignment = Alignment(
        horizontal="center",
    )

    worksheet.merge_cells("A4:H4")
    scope_label = (
        "Filtered Result"
        if normalized_scope == "filtered"
        else "Whole Live Stock"
    )
    worksheet["A4"] = (
        "LIVE STOCK REPORT"
        f" | Scope: {scope_label}"
        " | Generated: "
        + datetime.now().strftime("%d-%m-%Y %H:%M")
    )
    worksheet["A4"].font = Font(italic=True, size=10)
    worksheet["A4"].alignment = Alignment(
        horizontal="center",
    )

    headers = [
        "HSN",
        "Product",
        "Product Code",
        "Unit",
        "Current Stock",
        "Purchase Price",
        "Stock Value",
        "Status",
    ]

    header_row = 5

    for column_index, header in enumerate(
        headers,
        start=1,
    ):
        cell = worksheet.cell(
            row=header_row,
            column=column_index,
            value=header,
        )
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="2F67D8",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    thin = Side(
        style="thin",
        color="D9E2F0",
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )

    first_data_row = 6

    for row_index, item in enumerate(
        live_items,
        start=first_data_row,
    ):
        values = [
            item.hsn_code,
            item.product_name,
            item.product_code,
            item.unit,
            float(item.current_stock),
            float(item.purchase_price),
            float(item.stock_value),
            item.stock_status,
        ]

        for column_index, value in enumerate(
            values,
            start=1,
        ):
            cell = worksheet.cell(
                row=row_index,
                column=column_index,
                value=value,
            )
            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
            )

        worksheet.cell(
            row=row_index,
            column=5,
        ).number_format = '#,##0.00'

        worksheet.cell(
            row=row_index,
            column=6,
        ).number_format = '₹#,##0.00'

        worksheet.cell(
            row=row_index,
            column=7,
        ).number_format = '₹#,##0.00'

    total_row = (
        first_data_row
        +
        len(live_items)
        +
        1
    )

    worksheet.merge_cells(
        start_row=total_row,
        start_column=1,
        end_row=total_row,
        end_column=4,
    )

    worksheet.cell(
        row=total_row,
        column=1,
        value="TOTAL",
    )

    total_quantity = sum(
        (
            item.current_stock
            for item
            in live_items
        ),
        0,
    )

    total_value = sum(
        (
            item.stock_value
            for item
            in live_items
        ),
        0,
    )

    worksheet.cell(
        row=total_row,
        column=5,
        value=float(
            total_quantity
        ),
    ).number_format = '#,##0.00'

    worksheet.cell(
        row=total_row,
        column=7,
        value=float(
            total_value
        ),
    ).number_format = '₹#,##0.00'

    for column_index in range(
        1,
        9,
    ):
        cell = worksheet.cell(
            row=total_row,
            column=column_index,
        )
        cell.font = Font(
            bold=True,
        )
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="EAF1FF",
        )
        cell.border = border

    widths = {
        1: 16,
        2: 34,
        3: 18,
        4: 12,
        5: 16,
        6: 18,
        7: 18,
        8: 16,
    }

    for column_index, width in widths.items():
        worksheet.column_dimensions[
            get_column_letter(
                column_index
            )
        ].width = width

    if live_items:
        worksheet.auto_filter.ref = (
            f"A5:H{first_data_row + len(live_items) - 1}"
        )

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)

    timestamp = (
        datetime.now()
        .strftime(
            "%Y%m%d_%H%M%S"
        )
    )

    filename = (
        "castle_live_stock_"
        f"{normalized_scope}_"
        f"{timestamp}.xlsx"
    )

    return StreamingResponse(
        stream,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                f'attachment; filename="{filename}"'
        },
    )


@router.get(
    "/movements",
    response_model=StockMovementResponse,
)
def get_stock_movements(
    product_id: int | None = Query(
        default=None,
        ge=1,
        description=(
            "Filter movements by Product ID"
        ),
    ),
    movement_type: str | None = Query(
        default=None,
        description=(
            "Filter by Purchase, "
            "Shop Floor Issue, or "
            "Finished Goods Receipt"
        ),
    ),
    start_date: datetime | None = Query(
        default=None,
        description=(
            "Include movements from this date/time"
        ),
    ),
    end_date: datetime | None = Query(
        default=None,
        description=(
            "Include movements up to this date/time"
        ),
    ),
    db: Session = Depends(
        get_db
    ),
    current_user: User = Depends(
        require_permission(
            "stock.view"
        )
    ),
):
    return (
        StockMovementService
        .get_movements(
            db=db,
            product_id=product_id,
            movement_type=movement_type,
            start_date=start_date,
            end_date=end_date,
        )
    )
