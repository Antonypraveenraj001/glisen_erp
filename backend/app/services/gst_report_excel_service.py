from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
)
from openpyxl.utils import (
    get_column_letter,
)


# ================================================================
# COMMON HELPERS
# ================================================================

def money(
    value,
) -> float:

    return float(
        Decimal(
            str(
                value
                if value is not None
                else 0
            )
        )
    )


def auto_size_columns(
    sheet,
) -> None:

    for column_cells in sheet.columns:

        max_length = 0

        column_letter = (
            get_column_letter(
                column_cells[0].column
            )
        )

        for cell in column_cells:

            value = (
                ""
                if cell.value is None
                else str(cell.value)
            )

            max_length = max(
                max_length,
                len(value),
            )

        sheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            42,
        )


# ================================================================
# SALES GST EXCEL
#
# IMPORTANT:
# Existing GST API imports this class.
# Keep this class name unchanged.
# ================================================================

class GSTReportExcelService:

    @staticmethod
    def build_excel(
        report: dict,
    ) -> BytesIO:

        workbook = Workbook()

        summary_sheet = (
            workbook.active
        )

        summary_sheet.title = (
            "Sales GST Summary"
        )

        detail_sheet = (
            workbook.create_sheet(
                "Sales GST Details"
            )
        )


        # ========================================================
        # SUMMARY
        # ========================================================

        summary_sheet[
            "A1"
        ] = "Sales GST Report"

        summary_sheet[
            "A1"
        ].font = Font(
            bold=True,
            size=16,
        )

        summary_sheet.merge_cells(
            "A1:D1"
        )


        summary_sheet[
            "A3"
        ] = "Start Date"

        summary_sheet[
            "B3"
        ] = (
            report.get(
                "start_date"
            )
            or "All"
        )


        summary_sheet[
            "A4"
        ] = "End Date"

        summary_sheet[
            "B4"
        ] = (
            report.get(
                "end_date"
            )
            or "All"
        )


        summary_sheet[
            "A6"
        ] = "Invoice Count"

        summary_sheet[
            "B6"
        ] = report.get(
            "invoice_count",
            0,
        )


        summary_sheet[
            "A7"
        ] = "Credit Note Count"

        summary_sheet[
            "B7"
        ] = report.get(
            "credit_note_count",
            0,
        )


        summary_sheet[
            "A9"
        ] = "Category"

        summary_sheet[
            "B9"
        ] = "Gross"

        summary_sheet[
            "C9"
        ] = "Credit"

        summary_sheet[
            "D9"
        ] = "Net"


        for cell in (
            summary_sheet[9]
        ):
            cell.font = Font(
                bold=True
            )

            cell.alignment = (
                Alignment(
                    horizontal="center"
                )
            )


        summary_rows = [
            (
                "Taxable Amount",
                report.get(
                    "gross_taxable_amount",
                    0,
                ),
                report.get(
                    "credit_taxable_amount",
                    0,
                ),
                report.get(
                    "net_taxable_amount",
                    0,
                ),
            ),
            (
                "CGST",
                report.get(
                    "gross_cgst_amount",
                    0,
                ),
                report.get(
                    "credit_cgst_amount",
                    0,
                ),
                report.get(
                    "net_cgst_amount",
                    0,
                ),
            ),
            (
                "SGST",
                report.get(
                    "gross_sgst_amount",
                    0,
                ),
                report.get(
                    "credit_sgst_amount",
                    0,
                ),
                report.get(
                    "net_sgst_amount",
                    0,
                ),
            ),
            (
                "IGST",
                report.get(
                    "gross_igst_amount",
                    0,
                ),
                report.get(
                    "credit_igst_amount",
                    0,
                ),
                report.get(
                    "net_igst_amount",
                    0,
                ),
            ),
            (
                "GST Total",
                report.get(
                    "gross_tax_amount",
                    0,
                ),
                report.get(
                    "credit_tax_amount",
                    0,
                ),
                report.get(
                    "net_tax_amount",
                    0,
                ),
            ),
            (
                "Invoice Total",
                report.get(
                    "gross_invoice_total",
                    0,
                ),
                report.get(
                    "credit_note_total",
                    0,
                ),
                report.get(
                    "net_sales_total",
                    0,
                ),
            ),
        ]


        row_number = 10

        for (
            label,
            gross,
            credit,
            net,
        ) in summary_rows:

            summary_sheet.cell(
                row=row_number,
                column=1,
                value=label,
            )

            summary_sheet.cell(
                row=row_number,
                column=2,
                value=money(
                    gross
                ),
            )

            summary_sheet.cell(
                row=row_number,
                column=3,
                value=money(
                    credit
                ),
            )

            summary_sheet.cell(
                row=row_number,
                column=4,
                value=money(
                    net
                ),
            )

            row_number += 1


        for row in range(
            10,
            row_number,
        ):

            for column in range(
                2,
                5,
            ):

                summary_sheet.cell(
                    row=row,
                    column=column,
                ).number_format = (
                    "#,##0.00"
                )


        # ========================================================
        # DETAIL
        # ========================================================

        headers = [
            "Date",
            "Document Number",
            "Document Type",
            "Reference Invoice",
            "Customer",
            "GSTIN",
            "Taxable Amount",
            "CGST",
            "SGST",
            "IGST",
            "GST Total",
            "Grand Total",
        ]


        detail_sheet.append(
            headers
        )


        for cell in (
            detail_sheet[1]
        ):
            cell.font = Font(
                bold=True
            )

            cell.alignment = (
                Alignment(
                    horizontal="center"
                )
            )


        for item in report.get(
            "items",
            [],
        ):

            detail_sheet.append(
                [
                    item.get(
                        "document_date"
                    ),
                    item.get(
                        "document_number"
                    ),
                    item.get(
                        "document_type"
                    ),
                    item.get(
                        "reference_invoice_number"
                    ),
                    item.get(
                        "company_name"
                    ),
                    item.get(
                        "gst_number"
                    ),
                    money(
                        item.get(
                            "taxable_amount"
                        )
                    ),
                    money(
                        item.get(
                            "cgst_amount"
                        )
                    ),
                    money(
                        item.get(
                            "sgst_amount"
                        )
                    ),
                    money(
                        item.get(
                            "igst_amount"
                        )
                    ),
                    money(
                        item.get(
                            "tax_amount"
                        )
                    ),
                    money(
                        item.get(
                            "grand_total"
                        )
                    ),
                ]
            )


        for row in (
            detail_sheet.iter_rows(
                min_row=2,
                min_col=7,
                max_col=12,
            )
        ):

            for cell in row:
                cell.number_format = (
                    "#,##0.00"
                )


        detail_sheet.freeze_panes = (
            "A2"
        )


        auto_size_columns(
            summary_sheet
        )

        auto_size_columns(
            detail_sheet
        )


        # ========================================================
        # OUTPUT
        # ========================================================

        output = BytesIO()

        workbook.save(
            output
        )

        output.seek(
            0
        )

        return output


    # ============================================================
    # OPTIONAL COMPATIBILITY METHOD
    # ============================================================

    @staticmethod
    def build_purchase_excel(
        report: dict,
    ) -> BytesIO:

        return (
            PurchaseGSTReportExcelService
            .build_excel(
                report
            )
        )


# ================================================================
# PURCHASE GST EXCEL
# ================================================================

class PurchaseGSTReportExcelService:

    @staticmethod
    def build_excel(
        report: dict,
    ) -> BytesIO:

        workbook = Workbook()

        summary_sheet = (
            workbook.active
        )

        summary_sheet.title = (
            "Purchase GST Summary"
        )

        detail_sheet = (
            workbook.create_sheet(
                "Purchase GST Details"
            )
        )


        # ========================================================
        # SUMMARY
        # ========================================================

        summary_sheet[
            "A1"
        ] = "Purchase GST Report"

        summary_sheet[
            "A1"
        ].font = Font(
            bold=True,
            size=16,
        )

        summary_sheet.merge_cells(
            "A1:B1"
        )


        summary_sheet[
            "A3"
        ] = "Start Date"

        summary_sheet[
            "B3"
        ] = (
            report.get(
                "start_date"
            )
            or "All"
        )


        summary_sheet[
            "A4"
        ] = "End Date"

        summary_sheet[
            "B4"
        ] = (
            report.get(
                "end_date"
            )
            or "All"
        )


        summary_sheet[
            "A6"
        ] = "Purchase Bill Count"

        summary_sheet[
            "B6"
        ] = report.get(
            "purchase_bill_count",
            report.get(
                "document_count",
                0,
            ),
        )


        purchase_summary = [
            (
                "Taxable Amount",
                report.get(
                    "total_taxable_amount",
                    report.get(
                        "taxable_amount",
                        0,
                    ),
                ),
            ),
            (
                "CGST",
                report.get(
                    "total_cgst_amount",
                    report.get(
                        "cgst_amount",
                        0,
                    ),
                ),
            ),
            (
                "SGST",
                report.get(
                    "total_sgst_amount",
                    report.get(
                        "sgst_amount",
                        0,
                    ),
                ),
            ),
            (
                "IGST",
                report.get(
                    "total_igst_amount",
                    report.get(
                        "igst_amount",
                        0,
                    ),
                ),
            ),
            (
                "GST Total",
                report.get(
                    "total_tax_amount",
                    report.get(
                        "total_gst",
                        0,
                    ),
                ),
            ),
            (
                "Purchase Total",
                report.get(
                    "total_bill_amount",
                    report.get(
                        "grand_total",
                        0,
                    ),
                ),
            ),
        ]


        summary_sheet[
            "A8"
        ] = "Category"

        summary_sheet[
            "B8"
        ] = "Amount"


        for cell in (
            summary_sheet[8]
        ):
            cell.font = Font(
                bold=True
            )


        row_number = 9

        for (
            label,
            value,
        ) in purchase_summary:

            summary_sheet.cell(
                row=row_number,
                column=1,
                value=label,
            )

            summary_sheet.cell(
                row=row_number,
                column=2,
                value=money(
                    value
                ),
            )

            summary_sheet.cell(
                row=row_number,
                column=2,
            ).number_format = (
                "#,##0.00"
            )

            row_number += 1


        # ========================================================
        # DETAIL
        # ========================================================

        headers = [
            "Date",
            "Purchase Bill Number",
            "Supplier",
            "Supplier GSTIN",
            "Taxable Amount",
            "CGST",
            "SGST",
            "IGST",
            "GST Total",
            "Bill Total",
        ]


        detail_sheet.append(
            headers
        )


        for cell in (
            detail_sheet[1]
        ):
            cell.font = Font(
                bold=True
            )

            cell.alignment = (
                Alignment(
                    horizontal="center"
                )
            )


        for item in report.get(
            "items",
            [],
        ):

            detail_sheet.append(
                [
                    item.get(
                        "bill_date",
                        item.get(
                            "document_date"
                        ),
                    ),
                    item.get(
                        "bill_number",
                        item.get(
                            "document_number"
                        ),
                    ),
                    item.get(
                        "supplier_name",
                        item.get(
                            "company_name"
                        ),
                    ),
                    item.get(
                        "supplier_gst_number",
                        item.get(
                            "gst_number"
                        ),
                    ),
                    money(
                        item.get(
                            "taxable_amount",
                            item.get(
                                "subtotal"
                            ),
                        )
                    ),
                    money(
                        item.get(
                            "cgst_amount"
                        )
                    ),
                    money(
                        item.get(
                            "sgst_amount"
                        )
                    ),
                    money(
                        item.get(
                            "igst_amount"
                        )
                    ),
                    money(
                        item.get(
                            "tax_amount",
                            item.get(
                                "total_gst"
                            ),
                        )
                    ),
                    money(
                        item.get(
                            "grand_total"
                        )
                    ),
                ]
            )


        for row in (
            detail_sheet.iter_rows(
                min_row=2,
                min_col=5,
                max_col=10,
            )
        ):

            for cell in row:
                cell.number_format = (
                    "#,##0.00"
                )


        detail_sheet.freeze_panes = (
            "A2"
        )


        auto_size_columns(
            summary_sheet
        )

        auto_size_columns(
            detail_sheet
        )


        # ========================================================
        # OUTPUT
        # ========================================================

        output = BytesIO()

        workbook.save(
            output
        )

        output.seek(
            0
        )

        return output