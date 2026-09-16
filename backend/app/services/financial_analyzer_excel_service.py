from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
)
from openpyxl.utils import (
    get_column_letter,
)


class FinancialAnalyzerExcelService:

    # ========================================================
    # DECIMAL
    # ========================================================

    @staticmethod
    def decimal(
        value,
    ) -> Decimal:
        return Decimal(
            str(
                value
                if value is not None
                else 0
            )
        )

    # ========================================================
    # BUILD EXCEL
    # ========================================================

    @staticmethod
    def build_excel(
        analysis: dict,
    ) -> BytesIO:

        workbook = Workbook()

        summary_sheet = (
            workbook.active
        )

        summary_sheet.title = (
            "Financial Summary"
        )

        production_sheet = (
            workbook.create_sheet(
                "Production Cost"
            )
        )

        expense_sheet = (
            workbook.create_sheet(
                "Expense Breakdown"
            )
        )

        # ====================================================
        # FINANCIAL SUMMARY
        # ====================================================

        summary_sheet["A1"] = (
            "Glisen ERP - Financial Analyzer"
        )

        summary_sheet["A1"].font = Font(
            bold=True,
            size=16,
        )

        summary_sheet.merge_cells(
            "A1:D1"
        )

        summary_sheet["A3"] = (
            "Start Date"
        )

        summary_sheet["B3"] = (
            analysis[
                "start_date"
            ]
            if analysis[
                "start_date"
            ] is not None
            else "All"
        )

        summary_sheet["A4"] = (
            "End Date"
        )

        summary_sheet["B4"] = (
            analysis[
                "end_date"
            ]
            if analysis[
                "end_date"
            ] is not None
            else "All"
        )

        # ====================================================
        # RECORD COUNTS
        # ====================================================

        summary_sheet["A6"] = (
            "Record Summary"
        )

        summary_sheet["A6"].font = Font(
            bold=True,
            size=12,
        )

        record_rows = [
            (
                "Effective Invoices",
                analysis[
                    "invoice_count"
                ],
            ),
            (
                "Issued Credit Notes",
                analysis[
                    "credit_note_count"
                ],
            ),
            (
                "Finished Products",
                analysis[
                    "finished_product_count"
                ],
            ),
            (
                "Expense Records",
                analysis[
                    "expense_count"
                ],
            ),
        ]

        row_number = 7

        for label, value in record_rows:

            summary_sheet.cell(
                row=row_number,
                column=1,
                value=label,
            )

            summary_sheet.cell(
                row=row_number,
                column=2,
                value=value,
            )

            row_number += 1

        # ====================================================
        # SALES
        # ====================================================

        row_number += 1

        summary_sheet.cell(
            row=row_number,
            column=1,
            value="Sales Revenue",
        ).font = Font(
            bold=True,
            size=12,
        )

        row_number += 1

        sales_rows = [
            (
                "Gross Sales Revenue",
                analysis[
                    "gross_sales"
                ],
            ),
            (
                "Less: Credit Notes",
                -FinancialAnalyzerExcelService
                .decimal(
                    analysis[
                        "credit_notes"
                    ]
                ),
            ),
            (
                "Net Sales Revenue",
                analysis[
                    "net_sales"
                ],
            ),
        ]

        for label, amount in sales_rows:

            summary_sheet.cell(
                row=row_number,
                column=1,
                value=label,
            )

            summary_sheet.cell(
                row=row_number,
                column=2,
                value=float(
                    FinancialAnalyzerExcelService
                    .decimal(
                        amount
                    )
                ),
            )

            summary_sheet.cell(
                row=row_number,
                column=2,
            ).number_format = (
                "#,##0.00;[Red]-#,##0.00"
            )

            if (
                label
                == "Net Sales Revenue"
            ):
                summary_sheet.cell(
                    row=row_number,
                    column=1,
                ).font = Font(
                    bold=True
                )

                summary_sheet.cell(
                    row=row_number,
                    column=2,
                ).font = Font(
                    bold=True
                )

            row_number += 1

        # ====================================================
        # PRODUCTION COST
        # ====================================================

        row_number += 1

        summary_sheet.cell(
            row=row_number,
            column=1,
            value="Production Cost",
        ).font = Font(
            bold=True,
            size=12,
        )

        row_number += 1

        production_rows = [
            (
                "Actual Material Cost",
                analysis[
                    "actual_material_cost"
                ],
            ),
            (
                "Actual Operation Cost",
                analysis[
                    "actual_operation_cost"
                ],
            ),
            (
                "Total Production Cost",
                analysis[
                    "total_production_cost"
                ],
            ),
        ]

        for label, amount in production_rows:

            summary_sheet.cell(
                row=row_number,
                column=1,
                value=label,
            )

            summary_sheet.cell(
                row=row_number,
                column=2,
                value=float(
                    FinancialAnalyzerExcelService
                    .decimal(
                        amount
                    )
                ),
            )

            summary_sheet.cell(
                row=row_number,
                column=2,
            ).number_format = (
                "#,##0.00"
            )

            if (
                label
                == "Total Production Cost"
            ):
                summary_sheet.cell(
                    row=row_number,
                    column=1,
                ).font = Font(
                    bold=True
                )

                summary_sheet.cell(
                    row=row_number,
                    column=2,
                ).font = Font(
                    bold=True
                )

            row_number += 1

        # ====================================================
        # BUSINESS COST
        # ====================================================

        row_number += 1

        summary_sheet.cell(
            row=row_number,
            column=1,
            value="Business Cost & Profit",
        ).font = Font(
            bold=True,
            size=12,
        )

        row_number += 1

        business_rows = [
            (
                "Company Expenses",
                analysis[
                    "total_expenses"
                ],
            ),
            (
                "Total Business Cost",
                analysis[
                    "total_business_cost"
                ],
            ),
            (
                "Net Profit / Loss",
                analysis[
                    "net_profit"
                ],
            ),
        ]

        for label, amount in business_rows:

            summary_sheet.cell(
                row=row_number,
                column=1,
                value=label,
            )

            summary_sheet.cell(
                row=row_number,
                column=2,
                value=float(
                    FinancialAnalyzerExcelService
                    .decimal(
                        amount
                    )
                ),
            )

            summary_sheet.cell(
                row=row_number,
                column=2,
            ).number_format = (
                "#,##0.00;[Red]-#,##0.00"
            )

            if label in {
                "Total Business Cost",
                "Net Profit / Loss",
            }:
                summary_sheet.cell(
                    row=row_number,
                    column=1,
                ).font = Font(
                    bold=True
                )

                summary_sheet.cell(
                    row=row_number,
                    column=2,
                ).font = Font(
                    bold=True
                )

            row_number += 1

        # ====================================================
        # ACCOUNTING NOTE
        # ====================================================

        row_number += 1

        summary_sheet.cell(
            row=row_number,
            column=1,
            value="Revenue Basis",
        ).font = Font(
            bold=True
        )

        summary_sheet.cell(
            row=row_number,
            column=2,
            value=(
                "Sales revenue excludes GST."
            ),
        )

        row_number += 1

        summary_sheet.cell(
            row=row_number,
            column=1,
            value="Production Cost Basis",
        ).font = Font(
            bold=True
        )

        summary_sheet.cell(
            row=row_number,
            column=2,
            value=(
                "Actual Shop Floor Issues "
                "+ Production Operations"
            ),
        )

        # ====================================================
        # PRODUCTION COST SHEET
        # ====================================================

        production_sheet["A1"] = (
            "Production Cost Summary"
        )

        production_sheet["A1"].font = Font(
            bold=True,
            size=15,
        )

        production_sheet.merge_cells(
            "A1:C1"
        )

        production_headers = [
            "Production Metric",
            "Amount",
            "Notes",
        ]

        production_sheet.append(
            []
        )

        production_sheet.append(
            production_headers
        )

        for cell in production_sheet[3]:
            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center"
            )

        production_data = [
            (
                "Finished Products",
                analysis[
                    "finished_product_count"
                ],
                (
                    "Finished products received "
                    "during the selected period"
                ),
            ),
            (
                "Actual Material Cost",
                analysis[
                    "actual_material_cost"
                ],
                (
                    "Sum of Shop Floor Issue "
                    "actual costs"
                ),
            ),
            (
                "Actual Operation Cost",
                analysis[
                    "actual_operation_cost"
                ],
                (
                    "Sum of Production Operation "
                    "actual costs"
                ),
            ),
            (
                "Total Production Cost",
                analysis[
                    "total_production_cost"
                ],
                (
                    "Material Cost "
                    "+ Operation Cost"
                ),
            ),
        ]

        for (
            metric,
            amount,
            notes,
        ) in production_data:

            if (
                metric
                == "Finished Products"
            ):
                amount_value = int(
                    amount
                )
            else:
                amount_value = float(
                    FinancialAnalyzerExcelService
                    .decimal(
                        amount
                    )
                )

            production_sheet.append(
                [
                    metric,
                    amount_value,
                    notes,
                ]
            )

        for row in production_sheet.iter_rows(
            min_row=4,
            min_col=2,
            max_col=2,
        ):
            for cell in row:

                if (
                    cell.row != 4
                ):
                    cell.number_format = (
                        "#,##0.00"
                    )

        production_sheet.cell(
            row=7,
            column=1,
        ).font = Font(
            bold=True
        )

        production_sheet.cell(
            row=7,
            column=2,
        ).font = Font(
            bold=True
        )

        # ====================================================
        # EXPENSE BREAKDOWN
        # ====================================================

        expense_sheet["A1"] = (
            "Company Expense Breakdown"
        )

        expense_sheet["A1"].font = Font(
            bold=True,
            size=15,
        )

        expense_sheet.merge_cells(
            "A1:C1"
        )

        expense_sheet.append(
            []
        )

        expense_headers = [
            "Category",
            "Expense Count",
            "Amount",
        ]

        expense_sheet.append(
            expense_headers
        )

        for cell in expense_sheet[3]:
            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center"
            )

        for item in analysis[
            "expense_breakdown"
        ]:

            expense_sheet.append(
                [
                    item[
                        "category"
                    ],
                    item[
                        "count"
                    ],
                    float(
                        FinancialAnalyzerExcelService
                        .decimal(
                            item[
                                "amount"
                            ]
                        )
                    ),
                ]
            )

        if analysis[
            "expense_breakdown"
        ]:

            for row in expense_sheet.iter_rows(
                min_row=4,
                min_col=3,
                max_col=3,
            ):
                for cell in row:
                    cell.number_format = (
                        "#,##0.00"
                    )

        # ====================================================
        # TOTAL EXPENSES
        # ====================================================

        total_expense_row = (
            expense_sheet.max_row
            + 2
        )

        expense_sheet.cell(
            row=total_expense_row,
            column=1,
            value="Total Company Expenses",
        ).font = Font(
            bold=True
        )

        expense_sheet.cell(
            row=total_expense_row,
            column=3,
            value=float(
                FinancialAnalyzerExcelService
                .decimal(
                    analysis[
                        "total_expenses"
                    ]
                )
            ),
        ).font = Font(
            bold=True
        )

        expense_sheet.cell(
            row=total_expense_row,
            column=3,
        ).number_format = (
            "#,##0.00"
        )

        # ====================================================
        # COLUMN WIDTHS
        # ====================================================

        for sheet in (
            summary_sheet,
            production_sheet,
            expense_sheet,
        ):

            for column_cells in sheet.columns:

                max_length = 0

                column_letter = (
                    get_column_letter(
                        column_cells[
                            0
                        ].column
                    )
                )

                for cell in column_cells:

                    value = (
                        ""
                        if cell.value
                        is None
                        else str(
                            cell.value
                        )
                    )

                    max_length = max(
                        max_length,
                        len(
                            value
                        ),
                    )

                sheet.column_dimensions[
                    column_letter
                ].width = min(
                    max_length
                    + 3,
                    55,
                )

        # ====================================================
        # FREEZE PANES
        # ====================================================

        production_sheet.freeze_panes = (
            "A4"
        )

        expense_sheet.freeze_panes = (
            "A4"
        )

        # ====================================================
        # OUTPUT
        # ====================================================

        output = BytesIO()

        workbook.save(
            output
        )

        output.seek(
            0
        )

        return output